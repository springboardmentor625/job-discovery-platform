import pdfplumber
from docx import Document
import re


# =========================================================
# SKILLS DATABASE
# =========================================================

SKILLS = [
    # -------------------------
    # Programming Languages
    # -------------------------

    "Python",
    "Java",
    "C++",
    "C#",
    "C",
    "JavaScript",
    "TypeScript",
    "Go",
    "Golang",
    "Rust",
    "Ruby",
    "PHP",
    "Kotlin",
    "Swift",
    "Dart",
    "Scala",
    "R",
    "MATLAB",
    "Perl",
    "Bash",
    "Shell",

    # -------------------------
    # Frontend
    # -------------------------

    "HTML",
    "HTML5",
    "CSS",
    "CSS3",
    "SASS",
    "SCSS",
    "Bootstrap",
    "Tailwind",
    "Tailwind CSS",
    "Material UI",
    "MUI",

    "React",
    "React.js",
    "Next.js",
    "Angular",
    "AngularJS",
    "Vue",
    "Vue.js",
    "Nuxt.js",
    "Svelte",
    "Redux",
    "Redux Toolkit",
    "React Native",

    # -------------------------
    # Backend
    # -------------------------

    "Node.js",
    "Express",
    "Express.js",
    "Django",
    "Django REST Framework",
    "DRF",
    "Flask",
    "FastAPI",
    "Spring",
    "Spring Boot",
    "Spring MVC",
    "Laravel",
    "CodeIgniter",
    ".NET",
    "ASP.NET",
    "ASP.NET Core",
    "Ruby on Rails",

    # -------------------------
    # APIs
    # -------------------------

    "REST",
    "REST API",
    "RESTful API",
    "GraphQL",
    "SOAP",
    "WebSocket",
    "WebSockets",
    "API Development",
    "API Integration",
    "Microservices",

    # -------------------------
    # Databases
    # -------------------------

    "SQL",
    "PostgreSQL",
    "MySQL",
    "MariaDB",
    "SQLite",
    "Oracle",
    "Oracle Database",
    "Microsoft SQL Server",
    "SQL Server",

    "MongoDB",
    "Redis",
    "Cassandra",
    "DynamoDB",
    "Firebase",
    "Firestore",
    "Elasticsearch",

    # -------------------------
    # Cloud
    # -------------------------

    "AWS",
    "Amazon Web Services",
    "EC2",
    "S3",
    "Lambda",
    "RDS",
    "ECS",
    "EKS",
    "CloudFront",

    "Azure",
    "Microsoft Azure",

    "Google Cloud",
    "GCP",
    "Google Cloud Platform",

    "Heroku",
    "Vercel",
    "Netlify",

    # -------------------------
    # DevOps
    # -------------------------

    "Docker",
    "Docker Compose",
    "Kubernetes",
    "Jenkins",
    "GitHub Actions",
    "GitLab CI",
    "CI/CD",
    "Continuous Integration",
    "Continuous Deployment",

    "Terraform",
    "Ansible",
    "Nginx",
    "Apache",

    # -------------------------
    # Version Control
    # -------------------------

    "Git",
    "GitHub",
    "GitLab",
    "Bitbucket",

    # -------------------------
    # Data Science
    # -------------------------

    "Machine Learning",
    "Deep Learning",
    "Artificial Intelligence",
    "AI",

    "Data Science",
    "Data Analysis",
    "Data Analytics",

    "TensorFlow",
    "PyTorch",
    "Keras",

    "Pandas",
    "NumPy",
    "SciPy",
    "Scikit-learn",
    "Matplotlib",
    "Seaborn",

    # -------------------------
    # NLP
    # -------------------------

    "Natural Language Processing",
    "NLP",
    "NLTK",
    "SpaCy",
    "Transformers",
    "Hugging Face",

    # -------------------------
    # AI / LLM
    # -------------------------

    "Generative AI",
    "Generative Artificial Intelligence",
    "Large Language Models",
    "LLM",
    "OpenAI",
    "LangChain",
    "LangGraph",
    "RAG",
    "Retrieval Augmented Generation",
    "Prompt Engineering",

    # -------------------------
    # Testing
    # -------------------------

    "Unit Testing",
    "Integration Testing",
    "Automation Testing",
    "Test Automation",

    "PyTest",
    "pytest",
    "Jest",
    "Mocha",
    "Chai",
    "Selenium",
    "Cypress",
    "Playwright",

    # -------------------------
    # Security
    # -------------------------

    "Cybersecurity",
    "Information Security",
    "OAuth",
    "OAuth2",
    "JWT",
    "JSON Web Token",
    "Authentication",
    "Authorization",

    # -------------------------
    # Architecture
    # -------------------------

    "Object Oriented Programming",
    "OOP",
    "Data Structures",
    "Algorithms",
    "System Design",
    "Software Architecture",
    "Design Patterns",

    # -------------------------
    # Tools
    # -------------------------

    "Postman",
    "Swagger",
    "OpenAPI",
    "Jira",
    "Confluence",
    "Trello",

    # -------------------------
    # Operating Systems
    # -------------------------

    "Linux",
    "Ubuntu",
    "Windows",
    "MacOS",
    "Unix",

    # -------------------------
    # Messaging
    # -------------------------

    "RabbitMQ",
    "Apache Kafka",
    "Kafka",
    "Celery",
    "Redis Queue",

    # -------------------------
    # Mobile
    # -------------------------

    "Android",
    "Android Development",
    "iOS",
    "Flutter",
    "React Native",

    # -------------------------
    # Methodologies
    # -------------------------

    "Agile",
    "Scrum",
    "Kanban",
    "SDLC",
    "DevOps",

    # -------------------------
    # Other Common Skills
    # -------------------------

    "JSON",
    "XML",
    "YAML",
    "Linux Administration",
    "Networking",
    "TCP/IP",
    "Computer Networks",
    "Debugging",
    "Problem Solving",
    "Technical Documentation",
]


# =========================================================
# NORMALIZE TEXT
# =========================================================

def normalize_text(text):
    """
    Normalize resume text for better matching.
    """
    if not text:
        return ""

    # Replace unicode spaces and special dashes
    text = text.replace("\xa0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("–", "-")
    text = text.replace("—", "-")
    text = text.replace("•", " ")
    text = text.replace("·", " ")

    # Strip surrogate or unencodable symbols safely
    text = text.encode("utf-8", "ignore").decode("utf-8")

    # Normalize multiple spaces
    text = re.sub(r"[ \t]+", " ", text)

    return text


# =========================================================
# EXTRACT TEXT
# =========================================================

def extract_text(file_path):
    extension = file_path.split(".")[-1].lower()

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------
    if extension == "pdf":
        text_parts = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
        except Exception as e:
            print("PDF EXTRACTION ERROR:", repr(e))
            return ""

        return "\n".join(text_parts)

    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------
    elif extension == "docx":
        try:
            doc = Document(file_path)
            paragraphs = []

            # Extract from headers (often holds candidate name and contact info)
            for section in doc.sections:
                if section.header:
                    for hp in section.header.paragraphs:
                        htext = hp.text.strip()
                        if htext and htext not in paragraphs:
                            paragraphs.append(htext)

            # Normal paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    paragraphs.append(text)

            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = []
                    for cell in row.cells:
                        ctext = cell.text.strip()
                        if ctext and ctext not in row_cells:
                            row_cells.append(ctext)
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))

            return "\n".join(paragraphs)

        except Exception as e:
            print("DOCX EXTRACTION ERROR:", repr(e))
            return ""

    return ""


# =========================================================
# SKILL MATCHING
# =========================================================

def skill_pattern(skill):

    """
    Creates a safer regex pattern for a skill.

    Allows things such as:

        Node.js
        C++
        C#
        .NET
        CI/CD
        Django REST Framework
    """

    escaped = re.escape(
        skill.lower()
    )

    # Normal word-boundary matching
    return (
        r"(?<![a-zA-Z0-9])"
        + escaped
        + r"(?![a-zA-Z0-9])"
    )


# =========================================================
# EXTRACT SKILLS
# =========================================================

def extract_skills(text):

    if not text:

        return []

    normalized_text = normalize_text(text)

    lower_text = normalized_text.lower()

    found = []

    # -----------------------------------------------------
    # IMPORTANT:
    # Match longer skills first.
    #
    # Example:
    #
    # "Django REST Framework"
    #
    # should be detected before:
    #
    # "Django"
    # "REST"
    # -----------------------------------------------------

    sorted_skills = sorted(
        SKILLS,
        key=len,
        reverse=True
    )

    matches = []

    for skill in sorted_skills:

        pattern = skill_pattern(skill)

        match = re.search(
            pattern,
            lower_text
        )

        if match:

            matches.append(
                (
                    match.start(),
                    match.end(),
                    skill
                )
            )

    # -----------------------------------------------------
    # Sort by position in resume
    #
    # This means skills are returned roughly in
    # the same order they appear in the resume.
    # -----------------------------------------------------

    matches.sort(
        key=lambda item: item[0]
    )

    # -----------------------------------------------------
    # Remove duplicates / overlapping skills
    # -----------------------------------------------------

    occupied_ranges = []

    for start, end, skill in matches:

        overlap = False

        for existing_start, existing_end in occupied_ranges:

            if (
                start < existing_end
                and end > existing_start
            ):

                overlap = True
                break

        if overlap:
            continue

        found.append(skill)

        occupied_ranges.append(
            (start, end)
        )

    return found


# =========================================================
# EXTRACT EMAIL
# =========================================================

def extract_email(text):

    emails = re.findall(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    )

    if emails:

        return emails[0].strip()

    return ""


# =========================================================
# EXTRACT PHONE
# =========================================================

def extract_phone(text):

    # -----------------------------------------------------
    # Indian phone numbers
    # -----------------------------------------------------

    phones = re.findall(
        r"(?:\+91[\s-]?)?[6-9]\d{9}",
        text
    )

    if phones:

        return phones[0].strip()

    # -----------------------------------------------------
    # International / formatted numbers
    # -----------------------------------------------------

    phones = re.findall(
        r"\+?\d[\d\s().-]{8,18}\d",
        text
    )

    if phones:

        return phones[0].strip()

    return ""


# =========================================================
# EXTRACT NAME
# =========================================================

def extract_name(text):

    lines = [
        line.strip()
        for line in text.split("\n")
        if line.strip()
    ]

    for line in lines[:8]:

        # Skip email
        if "@" in line:
            continue

        # Skip lines containing numbers
        if re.search(r"\d", line):
            continue

        # Skip obvious headings
        if line.lower() in [
            "resume",
            "curriculum vitae",
            "cv",
            "profile",
            "summary",
        ]:
            continue

        words = line.split()

        if 2 <= len(words) <= 5:

            # Avoid very long lines
            if len(line) <= 60:

                return line

    return ""


# =========================================================
# EXTRACT EXPERIENCE
# =========================================================

def extract_experience(text):

    lower_text = text.lower()

    patterns = [

        r"(\d+(?:\.\d+)?\+?)\s+years?\s+(?:of\s+)?experience",

        r"(\d+(?:\.\d+)?\+?)\s+years?\s+(?:in|of)",

        r"(\d+(?:\.\d+)?\+?)\s+yrs?\s+(?:of\s+)?experience",

        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?\+?)\s*years?",

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            lower_text
        )

        if match:

            experience_value = (
                match.group(1)
            )

            return (
                f"{experience_value} Years"
            )

    if "fresher" in lower_text:

        return "Fresher"

    if "fresh graduate" in lower_text:

        return "Fresher"

    if "recent graduate" in lower_text:

        return "Fresher"

    if "intern" in lower_text:

        return "Intern"

    return ""


# =========================================================
# EXTRACT EDUCATION
# =========================================================

def extract_education(text):

    education_keywords = [

        "b.tech",
        "btech",
        "b.e",
        "be ",
        "bachelor",
        "b.sc",
        "bsc",
        "bca",
        "bba",

        "m.tech",
        "mtech",
        "m.e",
        "master",
        "m.sc",
        "msc",
        "mca",
        "mba",

        "phd",
        "doctorate",

        "computer science",
        "information technology",
        "software engineering",

    ]

    education_lines = []

    for line in text.split("\n"):

        clean_line = line.strip()

        if not clean_line:
            continue

        lower_line = clean_line.lower()

        if any(
            keyword in lower_line
            for keyword in education_keywords
        ):

            education_lines.append(
                clean_line
            )

    if education_lines:

        # Return all matching education lines
        return "\n".join(
            education_lines
        )

    return ""


# =========================================================
# EXTRACT PROJECTS
# =========================================================

def extract_projects(text):

    project_lines = []

    keywords = [

        "project",
        "developed",
        "built",
        "created",
        "implemented",
        "designed",
        "application",
        "platform",
        "system",

    ]

    for line in text.split("\n"):

        clean_line = line.strip()

        if not clean_line:
            continue

        lower_line = clean_line.lower()

        if any(
            keyword in lower_line
            for keyword in keywords
        ):

            project_lines.append(
                clean_line
            )

    return project_lines


# =========================================================
# EXTRACT CERTIFICATIONS
# =========================================================

def extract_certifications(text):

    certification_lines = []

    keywords = [

        "certificate",
        "certification",
        "certified",

    ]

    for line in text.split("\n"):

        clean_line = line.strip()

        if not clean_line:
            continue

        lower_line = clean_line.lower()

        if any(
            keyword in lower_line
            for keyword in keywords
        ):

            certification_lines.append(
                clean_line
            )

    return certification_lines


# =========================================================
# MAIN PARSER
# =========================================================

def parse_resume(file_path):

    print("=================================")
    print("PARSING RESUME")
    print("FILE:", file_path)
    print("=================================")

    # -----------------------------------------------------
    # TEXT
    # -----------------------------------------------------

    text = extract_text(file_path)

    if not text:

        print(
            "NO TEXT EXTRACTED FROM RESUME"
        )

        return {
            "text": "",
            "skills": [],
            "name": "",
            "email": "",
            "phone": "",
            "experience": "",
            "education": "",
            "projects": [],
            "certifications": [],
        }

    print("TEXT EXTRACTED")

    # -----------------------------------------------------
    # NORMALIZE
    # -----------------------------------------------------

    text = normalize_text(text)

    # -----------------------------------------------------
    # PARSE
    # -----------------------------------------------------

    skills = extract_skills(text)

    name = extract_name(text)

    email = extract_email(text)

    phone = extract_phone(text)

    experience = extract_experience(text)

    education = extract_education(text)

    projects = extract_projects(text)

    certifications = extract_certifications(text)

    # -----------------------------------------------------
    # RESULT
    # -----------------------------------------------------

    result = {

        "text": text,

        "skills": skills,

        "name": name,

        "email": email,

        "phone": phone,

        "experience": experience,

        "education": education,

        "projects": projects,

        "certifications": certifications,
    }

    # -----------------------------------------------------
    # DEBUG
    # -----------------------------------------------------

    print("=================================")
    print("PARSED RESUME")

    print(
        "SKILLS:",
        skills
    )

    print(
        "SKILL COUNT:",
        len(skills)
    )

    print(
        "EMAIL:",
        email
    )

    print(
        "PHONE:",
        phone
    )

    print(
        "EXPERIENCE:",
        experience
    )

    print(
        "EDUCATION:",
        education
    )

    print(
        "PROJECTS:",
        projects
    )

    print(
        "CERTIFICATIONS:",
        certifications
    )

    print("=================================")

    return result
