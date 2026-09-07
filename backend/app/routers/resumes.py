import io
import uuid
from pathlib import Path

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    UploadFile,
    status,
)

from fastapi.security import (
    HTTPBearer,
    HTTPAuthorizationCredentials,
)

from sqlalchemy.orm import Session

from docx import Document
from pypdf import PdfReader

from ..database import get_db

from ..models import (
    User,
    Resume,
)

from ..utils.security import verify_access_token

from ..services.resume_service import (
    parse_resume,
)


# =========================================================
# ROUTER
# =========================================================

router = APIRouter(
    prefix="/api/resumes",
    tags=["Resumes"],
)

security = HTTPBearer()


# =========================================================
# UPLOAD DIRECTORY
# =========================================================

BASE_DIR = Path(
    __file__
).resolve().parents[2]

UPLOAD_DIR = BASE_DIR / "uploads"

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# =========================================================
# GET CURRENT USER ID
# =========================================================

def get_current_user_id(
    credentials: HTTPAuthorizationCredentials,
):

    token = credentials.credentials

    payload = verify_access_token(
        token
    )

    if payload is None:

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid or expired token",
        )

    user_id = payload.get(
        "sub"
    )

    if user_id is None:

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid token",
        )

    try:

        return int(
            user_id
        )

    except (
        TypeError,
        ValueError
    ):

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid user ID in token",
        )


# =========================================================
# EXTRACT TEXT FROM PDF
# =========================================================

def extract_text_from_pdf(
    file_bytes: bytes,
) -> str:

    try:

        reader = PdfReader(
            io.BytesIO(
                file_bytes
            )
        )

        text_parts = []

        for page in reader.pages:

            page_text = (
                page.extract_text()
            )

            if page_text:

                text_parts.append(
                    page_text
                )

        return "\n".join(
            text_parts
        )

    except Exception as e:

        raise HTTPException(
            status_code=400,
            detail=(
                "Unable to read the PDF resume. "
                f"Error: {str(e)}"
            ),
        )


# =========================================================
# EXTRACT TEXT FROM DOCX
# =========================================================

def extract_text_from_docx(
    file_bytes: bytes,
) -> str:

    try:

        document = Document(
            io.BytesIO(
                file_bytes
            )
        )

        text_parts = []

        # -------------------------------------------------
        # PARAGRAPHS
        # -------------------------------------------------

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text_parts.append(
                    paragraph.text
                )

        # -------------------------------------------------
        # TABLES
        # -------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    if cell.text.strip():

                        row_text.append(
                            cell.text
                        )

                if row_text:

                    text_parts.append(
                        " ".join(
                            row_text
                        )
                    )

        return "\n".join(
            text_parts
        )

    except Exception as e:

        raise HTTPException(
            status_code=400,
            detail=(
                "Unable to read the DOCX resume. "
                f"Error: {str(e)}"
            ),
        )


# =========================================================
# EXTRACT RESUME TEXT
# =========================================================

def extract_resume_text(
    file_bytes: bytes,
    extension: str,
) -> str:

    extension = extension.lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_bytes
        )

    if extension == ".docx":

        return extract_text_from_docx(
            file_bytes
        )

    if extension == ".doc":

        raise HTTPException(
            status_code=400,
            detail=(
                "Old .doc files are not supported "
                "for automatic text extraction. "
                "Please upload the resume as PDF or DOCX."
            ),
        )

    raise HTTPException(
        status_code=400,
        detail=(
            "Unsupported resume format. "
            "Please upload PDF or DOCX."
        ),
    )


# =========================================================
# UPLOAD RESUME
# =========================================================

@router.post(
    "/upload",
    status_code=status.HTTP_201_CREATED,
)
async def upload_resume(

    file: UploadFile = File(...),

    credentials: HTTPAuthorizationCredentials = Depends(
        security
    ),

    db: Session = Depends(get_db),
):

    # -----------------------------------------------------
    # 1. GET USER
    # -----------------------------------------------------

    user_id = get_current_user_id(
        credentials
    )

    user = (
        db.query(User)
        .filter(
            User.user_id == user_id
        )
        .first()
    )

    if not user:

        raise HTTPException(
            status_code=404,
            detail="User not found",
        )

    # -----------------------------------------------------
    # 2. CHECK FILE NAME
    # -----------------------------------------------------

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="Resume file name is missing.",
        )

    original_filename = (
        file.filename
    )

    extension = Path(
        original_filename
    ).suffix.lower()

    allowed_extensions = {
        ".pdf",
        ".docx",
        ".doc",
    }

    if extension not in allowed_extensions:

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid file type. "
                "Please upload PDF, DOC, or DOCX."
            ),
        )

    # -----------------------------------------------------
    # 3. READ FILE
    # -----------------------------------------------------

    file_bytes = await file.read()

    if not file_bytes:

        raise HTTPException(
            status_code=400,
            detail="The uploaded resume is empty.",
        )

    # -----------------------------------------------------
    # 4. CHECK FILE SIZE
    # -----------------------------------------------------

    max_size = 5 * 1024 * 1024

    if len(file_bytes) > max_size:

        raise HTTPException(
            status_code=400,
            detail=(
                "Resume file size must not exceed 5 MB."
            ),
        )

    # -----------------------------------------------------
    # 5. EXTRACT RAW TEXT
    # -----------------------------------------------------

    resume_text = extract_resume_text(
        file_bytes,
        extension,
    )

    if not resume_text.strip():

        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text was found in the resume. "
                "Please upload a text-based PDF or DOCX file."
            ),
        )

    # -----------------------------------------------------
    # 6. PARSE RESUME
    # -----------------------------------------------------

    parsed_resume = parse_resume(
        resume_text
    )

    extracted_skills = (
        parsed_resume[
            "extracted_skills"
        ]
    )

    extracted_experience = (
        parsed_resume[
            "extracted_experience"
        ]
    )

    extracted_education = (
        parsed_resume[
            "extracted_education"
        ]
    )

    # -----------------------------------------------------
    # 7. MAKE UNIQUE FILE NAME
    # -----------------------------------------------------

    unique_filename = (
        f"{uuid.uuid4().hex}"
        f"{extension}"
    )

    file_path = (
        UPLOAD_DIR /
        unique_filename
    )

    # -----------------------------------------------------
    # 8. SAVE PHYSICAL FILE
    # -----------------------------------------------------

    try:

        with open(
            file_path,
            "wb",
        ) as output_file:

            output_file.write(
                file_bytes
            )

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to save resume file. "
                f"Error: {str(e)}"
            ),
        )

    # -----------------------------------------------------
    # 9. REMOVE DEFAULT FROM OLD RESUMES
    # -----------------------------------------------------

    existing_resumes = (
        db.query(Resume)
        .filter(
            Resume.user_id == user_id
        )
        .all()
    )

    for existing_resume in existing_resumes:

        existing_resume.is_default = False

    # -----------------------------------------------------
    # 10. CREATE RESUME RECORD
    # -----------------------------------------------------

    new_resume = Resume(

        user_id=user_id,

        resume_name=original_filename,

        file_path=str(
            Path("uploads")
            / unique_filename
        ),

        resume_text=(
            parsed_resume[
                "resume_text"
            ]
        ),

        extracted_skills=(
            extracted_skills
        ),

        extracted_experience=(
            extracted_experience
        ),

        extracted_education=(
            extracted_education
        ),

        is_default=True,
    )

    db.add(
        new_resume
    )

    # -----------------------------------------------------
    # 11. SAVE DATABASE
    # -----------------------------------------------------

    try:

        db.commit()

        db.refresh(
            new_resume
        )

    except Exception as e:

        db.rollback()

        try:

            if file_path.exists():

                file_path.unlink()

        except Exception:

            pass

        raise HTTPException(
            status_code=500,
            detail=(
                "Database error while saving resume. "
                f"Error: {str(e)}"
            ),
        )

    # -----------------------------------------------------
    # 12. RETURN RESPONSE
    # -----------------------------------------------------

    return {

        "message":
            "Resume uploaded and parsed successfully",

        "resume_id":
            new_resume.resume_id,

        "user_id":
            new_resume.user_id,

        "resume_name":
            new_resume.resume_name,

        "file_path":
            new_resume.file_path,

        "resume_text":
            new_resume.resume_text,

        "extracted_skills":
            new_resume.extracted_skills,

        "extracted_experience":
            new_resume.extracted_experience,

        "extracted_education":
            new_resume.extracted_education,

        "uploaded_at":
            new_resume.uploaded_at,

        "is_default":
            new_resume.is_default,
    }


# =========================================================
# GET CURRENT USER'S RESUMES
# =========================================================

@router.get("/me")
def get_my_resumes(

    credentials: HTTPAuthorizationCredentials = Depends(
        security
    ),

    db: Session = Depends(get_db),
):

    user_id = get_current_user_id(
        credentials
    )

    user = (
        db.query(User)
        .filter(
            User.user_id == user_id
        )
        .first()
    )

    if not user:

        raise HTTPException(
            status_code=404,
            detail="User not found",
        )

    resumes = (
        db.query(Resume)
        .filter(
            Resume.user_id == user_id
        )
        .order_by(
            Resume.uploaded_at.desc()
        )
        .all()
    )

    resume_list = []

    for resume in resumes:

        resume_list.append({

            "resume_id":
                resume.resume_id,

            "user_id":
                resume.user_id,

            "resume_name":
                resume.resume_name,

            "file_path":
                resume.file_path,

            "resume_text":
                resume.resume_text,

            "extracted_skills":
                resume.extracted_skills,

            "extracted_experience":
                resume.extracted_experience,

            "extracted_education":
                resume.extracted_education,

            "uploaded_at":
                resume.uploaded_at,

            "is_default":
                resume.is_default,
        })

    return {

        "user_id":
            user_id,

        "resumes":
            resume_list,

        "count":
            len(resume_list),
    }


# =========================================================
# DELETE RESUME
# =========================================================

@router.delete(
    "/{resume_id}"
)
def delete_resume(

    resume_id: int,

    credentials: HTTPAuthorizationCredentials = Depends(
        security
    ),

    db: Session = Depends(get_db),
):

    user_id = get_current_user_id(
        credentials
    )

    resume = (
        db.query(Resume)
        .filter(
            Resume.resume_id == resume_id,
            Resume.user_id == user_id,
        )
        .first()
    )

    if not resume:

        raise HTTPException(
            status_code=404,
            detail="Resume not found.",
        )

    was_default = (
        resume.is_default
    )

    old_file_path = (
        resume.file_path
    )

    # -----------------------------------------------------
    # DELETE DATABASE RECORD
    # -----------------------------------------------------

    try:

        db.delete(
            resume
        )

        db.commit()

    except Exception as e:

        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to delete resume. "
                f"Error: {str(e)}"
            ),
        )

    # -----------------------------------------------------
    # DELETE PHYSICAL FILE
    # -----------------------------------------------------

    if old_file_path:

        physical_path = (
            BASE_DIR /
            old_file_path
        )

        try:

            if physical_path.exists():

                physical_path.unlink()

        except Exception:

            pass

    # -----------------------------------------------------
    # SET NEW DEFAULT
    # -----------------------------------------------------

    if was_default:

        next_resume = (
            db.query(Resume)
            .filter(
                Resume.user_id == user_id
            )
            .order_by(
                Resume.uploaded_at.desc()
            )
            .first()
        )

        if next_resume:

            next_resume.is_default = True

            try:

                db.commit()

            except Exception:

                db.rollback()

    # -----------------------------------------------------
    # RETURN
    # -----------------------------------------------------

    return {

        "message":
            "Resume deleted successfully",

        "resume_id":
            resume_id,
    }
